import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pdf_engine import sanitize_and_compute_items, fmt_money

# Paleta corporativa oficial F3R
COLOR_BLACK = "000000"
COLOR_GOLD = "D4AF37"
COLOR_GOLD_LIGHT = "FEF3C7"
COLOR_GRAY_LIGHT = "F8FAFC"
COLOR_GRAY_BORDER = "CBD5E1"
COLOR_DARK_TEXT = "0F172A"
COLOR_GRAY_TEXT = "64748B"

def get_thin_border(color=COLOR_GRAY_BORDER):
    side = Side(border_style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)

def get_gold_border():
    side = Side(border_style="medium", color=COLOR_GOLD)
    return Border(left=side, right=side, top=side, bottom=side)

def build_quote_excel(evt):
    """
    Genera un archivo Excel (.xlsx) corporativo de alta gama para F3R Producciones S.A.S.
    Contiene dos hojas:
    1. 'Propuesta Comercial': Ficha para el cliente con precios de venta, fórmulas y condiciones.
    2. 'Matriz de Costos Interna': Ficha de análisis financiero con costos y utilidad ex-ante.
    """
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # HOJA 1: PROPUESTA COMERCIAL OFICIAL
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Propuesta Comercial"
    ws1.views.sheetView[0].showGridLines = True
    
    calc = sanitize_and_compute_items(evt.get("items", []))
    items = calc["items"]
    
    # Banner Superior Negro y Oro F3R
    ws1.merge_cells("A1:E2")
    cell_top = ws1["A1"]
    cell_top.value = "F3R PRODUCCIONES S.A.S. — PROPUESTA COMERCIAL"
    cell_top.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    cell_top.fill = PatternFill(start_color=COLOR_BLACK, end_color=COLOR_BLACK, fill_type="solid")
    cell_top.alignment = Alignment(horizontal="center", vertical="center")
    
    ws1.merge_cells("A3:E3")
    bar_gold = ws1["A3"]
    bar_gold.value = f"COTIZACIÓN OFICIAL: {evt.get('id', 'V1-1009')}  |  FECHA: {evt.get('created_at', datetime.now().strftime('%Y-%m-%d'))}  |  MONEDA: COP"
    bar_gold.font = Font(name="Arial", size=9, bold=True, color="000000")
    bar_gold.fill = PatternFill(start_color=COLOR_GOLD, end_color=COLOR_GOLD, fill_type="solid")
    bar_gold.alignment = Alignment(horizontal="center", vertical="center")
    
    # Bloque de Información del Cliente y Proveedor
    ws1["A5"] = "DATOS DEL CLIENTE"
    ws1["A5"].font = Font(name="Arial", size=10, bold=True, color=COLOR_GOLD)
    ws1["A5"].fill = PatternFill(start_color=COLOR_BLACK, end_color=COLOR_BLACK, fill_type="solid")
    ws1.merge_cells("A5:B5")
    
    ws1["C5"] = "EMPRESA EMISORA & CONDICIONES"
    ws1["C5"].font = Font(name="Arial", size=10, bold=True, color=COLOR_GOLD)
    ws1["C5"].fill = PatternFill(start_color=COLOR_BLACK, end_color=COLOR_BLACK, fill_type="solid")
    ws1.merge_cells("C5:E5")
    
    prov = evt.get("provider_info", {})
    forma_pago = evt.get("forma_pago") or prov.get("forma_pago", "CONTADO")
    
    client_rows = [
        ("Cliente / Razón Social:", evt.get("client", "Cliente")),
        ("NIT / Documento:", evt.get("nit") or "N/A"),
        ("Nombre del Evento:", evt.get("event_name", "Evento F3R")),
        ("Lugar / Locación:", evt.get("location", "Por definir")),
        ("Aforo Estimado:", f"{evt.get('pax', 0)} personas"),
        ("Fecha del Evento:", str(evt.get("event_date", "Por definir"))),
    ]
    
    prov_rows = [
        ("Razón Social:", prov.get("razon_social", "F3R PRODUCCIONES S.A.S.") + f" (NIT {prov.get('nit', '902106670-3')})"),
        ("Atención Comercial:", "Francisco Ricaurte — WhatsApp: +57 333 434 0735"),
        ("Correo / Instagram:", "f3rproducciones@gmail.com | @f3rproducciones"),
        ("Dirección:", prov.get("direccion", "Cartagena de Indias Diagonal 31 80B75 Of 501")),
        ("Forma de Pago:", f"{forma_pago}"),
        ("Cuenta Bancaria:", "Cuenta de Ahorros Bancolombia # 788 000 169-77"),
    ]
    
    for i in range(6):
        r = 6 + i
        ws1[f"A{r}"] = client_rows[i][0]
        ws1[f"A{r}"].font = Font(name="Arial", size=8.5, bold=True, color=COLOR_DARK_TEXT)
        ws1[f"B{r}"] = client_rows[i][1]
        ws1[f"B{r}"].font = Font(name="Arial", size=8.5, color=COLOR_DARK_TEXT)
        
        ws1[f"C{r}"] = prov_rows[i][0]
        ws1[f"C{r}"].font = Font(name="Arial", size=8.5, bold=True, color=COLOR_DARK_TEXT)
        ws1[f"D{r}"] = prov_rows[i][1]
        ws1[f"D{r}"].font = Font(name="Arial", size=8.5, color=COLOR_DARK_TEXT)
        ws1.merge_cells(f"D{r}:E{r}")
        
    current_row = 13
    
    # Encabezados de la Tabla de Servicios
    headers = ["Item", "Descripción del Servicio", "Cantidad", "Valor Unitario (COP)", "Valor Total (COP)"]
    for c_idx, h in enumerate(headers, start=1):
        col_letter = get_column_letter(c_idx)
        cell = ws1[f"{col_letter}{current_row}"]
        cell.value = h
        cell.font = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=COLOR_BLACK, end_color=COLOR_BLACK, fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3] else ("right" if c_idx in [4, 5] else "left"), vertical="center")
        cell.border = get_gold_border()
        
    current_row += 1
    
    # Agrupación por categorías
    grouped = {}
    for it in items:
        cat = it["categoria"]
        if cat not in grouped:
            grouped[cat] = []
        grouped[cat].append(it)
        
    subtotal_rows = []
    
    for cat_name, cat_items in grouped.items():
        # Fila Encabezado Categoría
        ws1.merge_cells(f"A{current_row}:E{current_row}")
        cat_cell = ws1[f"A{current_row}"]
        cat_cell.value = f"  CATEGORÍA: {cat_name.upper()}"
        cat_cell.font = Font(name="Arial", size=9, bold=True, color=COLOR_GOLD)
        cat_cell.fill = PatternFill(start_color=COLOR_BLACK, end_color=COLOR_BLACK, fill_type="solid")
        cat_cell.alignment = Alignment(horizontal="left", vertical="center")
        cat_start_row = current_row + 1
        current_row += 1
        
        for idx, it in enumerate(cat_items, start=1):
            r = current_row
            ws1[f"A{r}"] = idx
            ws1[f"A{r}"].alignment = Alignment(horizontal="center")
            ws1[f"B{r}"] = it["item"]
            ws1[f"C{r}"] = it["cantidad"]
            ws1[f"C{r}"].alignment = Alignment(horizontal="center")
            ws1[f"D{r}"] = it["precio_unit"]
            ws1[f"D{r}"].number_format = '"$"#,##0'
            ws1[f"E{r}"] = f"=C{r}*D{r}"
            ws1[f"E{r}"].number_format = '"$"#,##0'
            
            fill_zebra = PatternFill(start_color="F8FAFC" if idx % 2 == 0 else "FFFFFF", fill_type="solid")
            for col_idx in range(1, 6):
                c = ws1[f"{get_column_letter(col_idx)}{r}"]
                c.font = Font(name="Arial", size=8.5, color=COLOR_DARK_TEXT)
                c.border = get_thin_border()
                c.fill = fill_zebra
                
            current_row += 1
            
        cat_end_row = current_row - 1
        
        # Fila Subtotal Categoría
        ws1.merge_cells(f"A{current_row}:D{current_row}")
        sub_lbl = ws1[f"A{current_row}"]
        sub_lbl.value = f"Subtotal {cat_name}:"
        sub_lbl.font = Font(name="Arial", size=9, bold=True, color=COLOR_DARK_TEXT)
        sub_lbl.alignment = Alignment(horizontal="right", vertical="center")
        sub_lbl.fill = PatternFill(start_color="F1F5F9", fill_type="solid")
        
        sub_val = ws1[f"E{current_row}"]
        sub_val.value = f"=SUM(E{cat_start_row}:E{cat_end_row})"
        sub_val.font = Font(name="Arial", size=9, bold=True, color=COLOR_DARK_TEXT)
        sub_val.number_format = '"$"#,##0'
        sub_val.fill = PatternFill(start_color="F1F5F9", fill_type="solid")
        
        for col_idx in range(1, 6):
            ws1[f"{get_column_letter(col_idx)}{current_row}"].border = get_thin_border()
            
        subtotal_rows.append(current_row)
        current_row += 1
        
    current_row += 1
    
    # Cuadros de Totales Oficiales
    formula_sum = "+".join([f"E{sr}" for sr in subtotal_rows]) if subtotal_rows else "0"
    
    ws1.merge_cells(f"C{current_row}:D{current_row}")
    ws1[f"C{current_row}"] = "SUBTOTAL SERVICIOS:"
    ws1[f"C{current_row}"].font = Font(name="Arial", size=10, bold=True, color=COLOR_DARK_TEXT)
    ws1[f"C{current_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws1[f"C{current_row}"].fill = PatternFill(start_color="F1F5F9", fill_type="solid")
    
    ws1[f"E{current_row}"] = f"={formula_sum}"
    ws1[f"E{current_row}"].font = Font(name="Arial", size=10, bold=True, color=COLOR_DARK_TEXT)
    ws1[f"E{current_row}"].number_format = '"$"#,##0'
    ws1[f"E{current_row}"].fill = PatternFill(start_color="F1F5F9", fill_type="solid")
    for col_idx in [3, 4, 5]:
        ws1[f"{get_column_letter(col_idx)}{current_row}"].border = get_thin_border()
    current_row += 1
    
    # Total General F3R en Oro y Negro
    ws1.merge_cells(f"C{current_row}:D{current_row}")
    ws1[f"C{current_row}"] = "VALOR TOTAL PROPUESTA:"
    ws1[f"C{current_row}"].font = Font(name="Arial", size=11, bold=True, color="000000")
    ws1[f"C{current_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws1[f"C{current_row}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    
    ws1[f"E{current_row}"] = f"=E{current_row - 1}"
    ws1[f"E{current_row}"].font = Font(name="Arial", size=12, bold=True, color="000000")
    ws1[f"E{current_row}"].number_format = '"$"#,##0'
    ws1[f"E{current_row}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    for col_idx in [3, 4, 5]:
        ws1[f"{get_column_letter(col_idx)}{current_row}"].border = get_gold_border()
    current_row += 2
    
    # Condiciones Comerciales
    ws1.merge_cells(f"A{current_row}:E{current_row}")
    ws1[f"A{current_row}"] = "CONDICIONES COMERCIALES Y DE CONTRATACIÓN F3R PRODUCCIONES S.A.S."
    ws1[f"A{current_row}"].font = Font(name="Arial", size=9, bold=True, color=COLOR_GOLD)
    ws1[f"A{current_row}"].fill = PatternFill(start_color=COLOR_BLACK, fill_type="solid")
    current_row += 1
    
    terminos_text = [
        "1. Vigencia: Esta propuesta tiene la vigencia indicada en la cotización y está sujeta a disponibilidad de fecha, equipos y personal.",
        "2. Reserva: La fecha se garantiza únicamente con la aceptación de la propuesta y el pago del anticipo acordado.",
        "3. Pagos: Los pagos deberán realizarse en las fechas establecidas. El incumplimiento podrá generar suspensión del servicio y/o cobro de intereses de mora.",
        "4. Cancelaciones: En caso de cancelación por parte del cliente, el anticipo no será reembolsable por las reservas y compromisos de F3R.",
        "5. Reprogramaciones: Todo cambio de fecha estará sujeto a disponibilidad y podrá generar costos adicionales.",
        "6. Servicios adicionales: Cualquier modificación, ampliación de horario o requerimiento extra será cotizado por separado.",
        "7. Montaje y logística: Permisos, transporte no estipulado, parqueaderos y gastos extraordinarios serán asumidos por el cliente.",
        "8. Equipos: El cliente será responsable por daños o pérdidas ocasionados por causas atribuibles al cliente o terceros bajo su responsabilidad.",
        "9. Lugar del evento: El cliente garantizará las condiciones de acceso, seguridad, espacio, energía y requerimientos de montaje.",
        "10. Fuerza mayor: F3R no será responsable por retrasos derivados de hechos imprevisibles fuera de su control.",
        "11. Responsabilidad: F3R responderá exclusivamente por los servicios expresamente contratados en esta propuesta.",
        "12. Aceptación: La aprobación de la cotización, firma, pago del anticipo o confirmación escrita implica la aceptación de estas condiciones."
    ]
    for t in terminos_text:
        ws1.merge_cells(f"A{current_row}:E{current_row}")
        ws1[f"A{current_row}"] = t
        ws1[f"A{current_row}"].font = Font(name="Arial", size=7.5, color=COLOR_DARK_TEXT)
        current_row += 1
        
    current_row += 2
    # Firmas
    ws1.merge_cells(f"A{current_row}:B{current_row}")
    ws1[f"A{current_row}"] = "Francisco Ricaurte"
    ws1[f"A{current_row}"].font = Font(name="Arial", size=9, bold=True, color=COLOR_DARK_TEXT)
    ws1[f"A{current_row}"].alignment = Alignment(horizontal="center")
    
    ws1.merge_cells(f"D{current_row}:E{current_row}")
    ws1[f"D{current_row}"] = f"ACEPTADO: {evt.get('client', 'CLIENTE')[:30]}"
    ws1[f"D{current_row}"].font = Font(name="Arial", size=9, bold=True, color=COLOR_DARK_TEXT)
    ws1[f"D{current_row}"].alignment = Alignment(horizontal="center")
    current_row += 1
    
    ws1.merge_cells(f"A{current_row}:B{current_row}")
    ws1[f"A{current_row}"] = "Gerente — F3R Producciones S.A.S.\nWhatsApp +57 333 434 0735\nf3rproducciones@gmail.com"
    ws1[f"A{current_row}"].font = Font(name="Arial", size=8, color=COLOR_GRAY_TEXT)
    ws1[f"A{current_row}"].alignment = Alignment(horizontal="center", wrap_text=True)
    
    ws1.merge_cells(f"D{current_row}:E{current_row}")
    ws1[f"D{current_row}"] = f"NIT / C.C.: {evt.get('nit', '')}\nFirma Autorizada / Sello de Aceptación"
    ws1[f"D{current_row}"].font = Font(name="Arial", size=8, color=COLOR_GRAY_TEXT)
    ws1[f"D{current_row}"].alignment = Alignment(horizontal="center", wrap_text=True)
    
    # Ajuste de anchos de columna
    ws1.column_dimensions["A"].width = 8
    ws1.column_dimensions["B"].width = 50
    ws1.column_dimensions["C"].width = 14
    ws1.column_dimensions["D"].width = 24
    ws1.column_dimensions["E"].width = 24

    # -------------------------------------------------------------
    # HOJA 2: MATRIZ DE COSTOS INTERNA F3R
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Matriz de Costos Interna")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:I2")
    c_mat = ws2["A1"]
    c_mat.value = f"MATRIZ DE COSTOS Y MARGEN INTERNO — {evt.get('id', 'V1-1009')} ({evt.get('client', '')})"
    c_mat.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    c_mat.fill = PatternFill(start_color=COLOR_BLACK, end_color=COLOR_BLACK, fill_type="solid")
    c_mat.alignment = Alignment(horizontal="center", vertical="center")
    
    mat_headers = ["Item", "Categoría", "Descripción del Servicio", "Cant", "Costo Unit (COP)", "Costo Total (COP)", "Venta Unit (COP)", "Venta Total (COP)", "Utilidad Neta"]
    for c_i, h in enumerate(mat_headers, start=1):
        col_l = get_column_letter(c_i)
        cell = ws2[f"{col_l}4"]
        cell.value = h
        cell.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if c_i in [1, 4] else ("right" if c_i >= 5 else "left"))
        cell.border = get_thin_border()
        
    m_row = 5
    for idx, it in enumerate(items, start=1):
        ws2[f"A{m_row}"] = idx
        ws2[f"B{m_row}"] = it["categoria"]
        ws2[f"C{m_row}"] = it["item"]
        ws2[f"D{m_row}"] = it["cantidad"]
        ws2[f"E{m_row}"] = it["costo_unit"]
        ws2[f"E{m_row}"].number_format = '"$"#,##0'
        ws2[f"F{m_row}"] = f"=D{m_row}*E{m_row}"
        ws2[f"F{m_row}"].number_format = '"$"#,##0'
        ws2[f"G{m_row}"] = it["precio_unit"]
        ws2[f"G{m_row}"].number_format = '"$"#,##0'
        ws2[f"H{m_row}"] = f"=D{m_row}*G{m_row}"
        ws2[f"H{m_row}"].number_format = '"$"#,##0'
        ws2[f"I{m_row}"] = f"=H{m_row}-F{m_row}"
        ws2[f"I{m_row}"].number_format = '"$"#,##0'
        
        for col_i in range(1, 10):
            ws2[f"{get_column_letter(col_i)}{m_row}"].border = get_thin_border()
            ws2[f"{get_column_letter(col_i)}{m_row}"].font = Font(name="Arial", size=8.5)
        m_row += 1
        
    # Totales Matriz
    ws2.merge_cells(f"A{m_row}:E{m_row}")
    ws2[f"A{m_row}"] = "TOTALES CONSOLIDADOS PRE-MONTAJE:"
    ws2[f"A{m_row}"].font = Font(name="Arial", size=10, bold=True, color="000000")
    ws2[f"A{m_row}"].alignment = Alignment(horizontal="right")
    ws2[f"A{m_row}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    
    ws2[f"F{m_row}"] = f"=SUM(F5:F{m_row-1})"
    ws2[f"F{m_row}"].font = Font(name="Arial", size=10, bold=True)
    ws2[f"F{m_row}"].number_format = '"$"#,##0'
    ws2[f"F{m_row}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    
    ws2[f"G{m_row}"] = "-"
    ws2[f"G{m_row}"].alignment = Alignment(horizontal="center")
    ws2[f"G{m_row}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    
    ws2[f"H{m_row}"] = f"=SUM(H5:H{m_row-1})"
    ws2[f"H{m_row}"].font = Font(name="Arial", size=10, bold=True)
    ws2[f"H{m_row}"].number_format = '"$"#,##0'
    ws2[f"H{m_row}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    
    ws2[f"I{m_row}"] = f"=H{m_row}-F{m_row}"
    ws2[f"I{m_row}"].font = Font(name="Arial", size=10, bold=True, color="059669")
    ws2[f"I{m_row}"].number_format = '"$"#,##0'
    ws2[f"I{m_row}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    
    for col_i in range(1, 10):
        ws2[f"{get_column_letter(col_i)}{m_row}"].border = get_gold_border()
        
    ws2.column_dimensions["A"].width = 7
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 45
    ws2.column_dimensions["D"].width = 8
    ws2.column_dimensions["E"].width = 18
    ws2.column_dimensions["F"].width = 20
    ws2.column_dimensions["G"].width = 18
    ws2.column_dimensions["H"].width = 20
    ws2.column_dimensions["I"].width = 20
    
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def build_consolidado_excel(events):
    """
    Genera el libro de consolidado comercial del CRM para F3R Producciones S.A.S.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado Comercial F3R"
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells("A1:K2")
    ws["A1"] = "F3R PRODUCCIONES S.A.S. — CONSOLIDADO COMERCIAL & CRM"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=COLOR_BLACK, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("A3:K3")
    ws["A3"] = f"REPORTE GENERAL GENERADO EL: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A3"].font = Font(name="Arial", size=9, bold=True, color="000000")
    ws["A3"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    
    cols = ["Código", "Cliente", "NIT", "Nombre Evento", "Fecha Evento", "Estado CRM", "Venta Total", "Costo Previsto", "Utilidad Prevista", "Costo Real", "Utilidad Real"]
    for i, c in enumerate(cols, start=1):
        cell = ws[f"{get_column_letter(i)}5"]
        cell.value = c
        cell.font = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E293B", fill_type="solid")
        cell.alignment = Alignment(horizontal="center" if i in [1, 5, 6] else ("right" if i >= 7 else "left"))
        cell.border = get_thin_border()
        
    r = 6
    for e in events:
        v_tot = sum(i["cantidad"] * i["precio_unit"] for i in e.get("items", []))
        c_tot = sum(i["cantidad"] * i.get("costo_unit", 0) for i in e.get("items", []))
        c_real = sum(c["costo_real"] for c in e.get("actual_costs", []))
        u_real = v_tot - c_real if c_real > 0 else (v_tot - c_tot)
        
        ws[f"A{r}"] = e.get("id", "")
        ws[f"B{r}"] = e.get("client", "")
        ws[f"C{r}"] = e.get("nit", "")
        ws[f"D{r}"] = e.get("event_name", "")
        ws[f"E{r}"] = str(e.get("event_date", ""))
        ws[f"F{r}"] = e.get("stage", "Cotizado")
        ws[f"G{r}"] = v_tot
        ws[f"G{r}"].number_format = '"$"#,##0'
        ws[f"H{r}"] = c_tot
        ws[f"H{r}"].number_format = '"$"#,##0'
        ws[f"I{r}"] = v_tot - c_tot
        ws[f"I{r}"].number_format = '"$"#,##0'
        ws[f"J{r}"] = c_real
        ws[f"J{r}"].number_format = '"$"#,##0'
        ws[f"K{r}"] = u_real
        ws[f"K{r}"].number_format = '"$"#,##0'
        
        for col_i in range(1, 12):
            ws[f"{get_column_letter(col_i)}{r}"].border = get_thin_border()
            ws[f"{get_column_letter(col_i)}{r}"].font = Font(name="Arial", size=8.5)
        r += 1
        
    # Totales
    ws.merge_cells(f"A{r}:F{r}")
    ws[f"A{r}"] = "TOTALES DE CARTERA COMERCIAL:"
    ws[f"A{r}"].font = Font(name="Arial", size=10, bold=True)
    ws[f"A{r}"].alignment = Alignment(horizontal="right")
    ws[f"A{r}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
    
    for c_i, letter in enumerate(["G", "H", "I", "J", "K"], start=7):
        ws[f"{letter}{r}"] = f"=SUM({letter}6:{letter}{r-1})"
        ws[f"{letter}{r}"].font = Font(name="Arial", size=10, bold=True)
        ws[f"{letter}{r}"].number_format = '"$"#,##0'
        ws[f"{letter}{r}"].fill = PatternFill(start_color=COLOR_GOLD, fill_type="solid")
        
    for col_i in range(1, 12):
        ws[f"{get_column_letter(col_i)}{r}"].border = get_gold_border()
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
